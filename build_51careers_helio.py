#!/usr/bin/env python3
"""Build 51Careers_AI_Investor_Overview.pptx in the Heliospace deck format.

Same investor narrative as 51Careers_Addition.pptx: Garamond, white canvas,
pale cards, numbered markers, logo on every slide.
Financing: RMB 20M at RMB 180M pre / RMB 200M post (10% equity).
"""

from __future__ import annotations

import os

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "51Careers_AI_Investor_Overview.pptx")
LOGO_DARK = os.path.join(ROOT, "assets", "51careers_logo_dark.png")
LOGO_LIGHT = os.path.join(ROOT, "assets", "51careers_logo_light.png")
COVER_BG = os.path.join(ROOT, "assets", "cover_bg.jpg")

# Heliospace visual system
BLUE = RGBColor(0x00, 0x7B, 0xCB)
BLUE_LT = RGBColor(0x6F, 0xBC, 0xE8)
BLUE_SOFT = RGBColor(0x9F, 0xC5, 0xDE)
TITLE = RGBColor(0x10, 0x22, 0x2F)
BODY = RGBColor(0x1E, 0x2A, 0x33)
MUTED = RGBColor(0x5C, 0x6B, 0x76)
PALE = RGBColor(0xF3, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC6, 0xD2, 0xDA)
COVER_CARD = RGBColor(0x0E, 0x22, 0x33)
COVER_LINE = RGBColor(0x3C, 0x5A, 0x72)
COVER_SUB = RGBColor(0xDC, 0xE8, 0xF1)
COVER_META = RGBColor(0xB9, 0xC7, 0xD2)
FONT = "Garamond"

SW, SH = Inches(13.333), Inches(7.5)
ML = Inches(0.90)
CW = Inches(11.53)


def apply_typeface(run, name):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def set_run(run, text, font=FONT, size=11, bold=False, color=BODY, italic=False):
    run.text = text
    apply_typeface(run, font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_text(shape, paragraphs, valign=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        shape.text_frame._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(valign, "t"),
        )
    except Exception:
        pass
    first = True
    for spec in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        if "space_after" in spec:
            p.space_after = Pt(spec["space_after"])
        if "space_before" in spec:
            p.space_before = Pt(spec["space_before"])
        run = p.add_run()
        set_run(
            run, spec["text"], font=spec.get("font", FONT), size=spec.get("size", 11),
            bold=spec.get("bold", False), color=spec.get("color", BODY), italic=spec.get("italic", False),
        )


def fill_solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    try:
        shape.line.fill.background()
    except (AttributeError, TypeError, ValueError):
        pass


def rect(slide, l, t, w, h, fill=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    if fill is None:
        s.fill.background()
        s.line.fill.background()
    else:
        fill_solid(s, fill)
    return s


def round_rect(slide, l, t, w, h, fill=PALE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill_solid(s, fill)
    try:
        s.adjustments[0] = 0.04
    except Exception:
        pass
    return s


def outlined_rect(slide, l, t, w, h, fill, line_color, line_pt=0.75):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line_color
    s.line.width = Pt(line_pt)
    try:
        s.adjustments[0] = 0.04
    except Exception:
        pass
    return s


def textbox(slide, l, t, w, h, paragraphs, valign=MSO_ANCHOR.TOP):
    s = slide.shapes.add_textbox(l, t, w, h)
    add_text(s, paragraphs, valign=valign)
    return s


def set_bg(slide, hex_color="FFFFFF"):
    cSld = slide._element.find(qn("p:cSld"))
    existing = cSld.find(qn("p:bg"))
    if existing is not None:
        cSld.remove(existing)
    bg = etree.Element(qn("p:bg"))
    bgPr = etree.SubElement(bg, qn("p:bgPr"))
    solid = etree.SubElement(bgPr, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", hex_color)
    etree.SubElement(bgPr, qn("a:effectLst"))
    cSld.find(qn("p:spTree")).addprevious(bg)


def logo(slide, light=False, cover=False):
    path = LOGO_LIGHT if light else LOGO_DARK
    if cover:
        slide.shapes.add_picture(path, Inches(0.90), Inches(0.75), Inches(2.75), Inches(0.78))
    else:
        slide.shapes.add_picture(path, Inches(11.00), Inches(0.52), Inches(1.45), Inches(0.41))


def chrome(slide, page: int):
    logo(slide, light=False, cover=False)
    textbox(slide, ML, Inches(7.10), Inches(8.50), Inches(0.26),
            [{"text": "51 Careers.AI   |   Proprietary and Confidential", "size": 8.5, "color": MUTED}])
    textbox(slide, Inches(12.10), Inches(7.10), Inches(0.35), Inches(0.26),
            [{"text": str(page), "size": 8.5, "color": MUTED, "align": PP_ALIGN.RIGHT}])


def header(slide, eyebrow: str, title: str, subtitle: str | None = None):
    textbox(slide, ML, Inches(0.50), Inches(8.50), Inches(0.30),
            [{"text": eyebrow.upper(), "size": 11, "bold": True, "color": BLUE}])
    title_size = 24 if len(title) > 42 else 29
    textbox(slide, ML, Inches(0.78), Inches(9.80), Inches(0.62),
            [{"text": title, "size": title_size, "bold": True, "color": TITLE}])
    if subtitle:
        textbox(slide, ML, Inches(1.42), CW, Inches(0.36),
                [{"text": subtitle, "size": 13.5, "color": BODY}])


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, "FFFFFF")
    return slide


def badge(slide, l, t, n, size=0.38):
    rect(slide, l, t, Inches(size), Inches(size), fill=BLUE)
    textbox(
        slide, l, t, Inches(size), Inches(size),
        [{"text": str(n), "size": 12, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}],
        valign=MSO_ANCHOR.MIDDLE,
    )


def num_block(slide, x, y, n, title, body, tw=5.60, bh=0.70):
    badge(slide, x, y, n)
    textbox(slide, x + Inches(0.55), y - Inches(0.02), Inches(tw), Inches(0.30),
            [{"text": title, "size": 15, "bold": True, "color": TITLE}])
    textbox(slide, x + Inches(0.55), y + Inches(0.32), Inches(tw), Inches(bh),
            [{"text": body, "size": 13.5, "color": BODY}])


def conclusion(slide, kicker: str, line: str):
    """Full-width closer so the canvas is used through the footer."""
    round_rect(slide, ML, Inches(6.18), CW, Inches(0.76), fill=PALE)
    textbox(slide, Inches(1.10), Inches(6.24), CW - Inches(0.40), Inches(0.22),
            [{"text": kicker.upper(), "size": 11, "bold": True, "color": BLUE}])
    textbox(slide, Inches(1.10), Inches(6.46), CW - Inches(0.40), Inches(0.40),
            [{"text": line, "size": 14, "bold": True, "color": TITLE}])


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def s_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(COVER_BG, Inches(0), Inches(0), SW, SH)
    logo(s, light=True, cover=True)
    textbox(s, ML, Inches(1.70), Inches(11.0), Inches(0.28),
            [{"text": "CONFIDENTIAL INVESTOR PRESENTATION", "size": 12, "bold": True, "color": BLUE_LT}])
    textbox(s, ML, Inches(2.00), Inches(11.20), Inches(0.78),
            [{"text": "51 Careers.AI", "size": 42, "bold": True, "color": WHITE}])
    textbox(s, ML, Inches(2.78), Inches(11.20), Inches(0.40),
            [{"text": "Building the global AI career platform — from resume to opportunity.",
              "size": 18, "color": COVER_SUB}])
    textbox(s, ML, Inches(3.22), Inches(11.20), Inches(0.28),
            [{"text": "AI RESUME   ×   AI CAREER   ×   GLOBAL TALENT ECOSYSTEM",
              "size": 13, "color": BLUE_SOFT}])
    cards = [
        ("PRE-MONEY", "RMB 180M", "≈ US$26.7 million"),
        ("RAISE  ·  10% EQUITY", "RMB 20M", "≈ US$2.97 million"),
        ("POST-MONEY", "RMB 200M", "≈ US$29.7 million"),
    ]
    for i, (lab, big, usd) in enumerate(cards):
        x = ML + i * Inches(3.91)
        outlined_rect(s, x, Inches(3.70), Inches(3.71), Inches(1.95), COVER_CARD, COVER_LINE, 0.75)
        textbox(s, x + Inches(0.22), Inches(3.88), Inches(3.27), Inches(0.28),
                [{"text": lab, "size": 12, "bold": True, "color": BLUE_LT}])
        textbox(s, x + Inches(0.22), Inches(4.20), Inches(3.27), Inches(0.70),
                [{"text": big, "size": 28, "bold": True, "color": WHITE}])
        textbox(s, x + Inches(0.22), Inches(4.92), Inches(3.27), Inches(0.42),
                [{"text": usd, "size": 14, "color": COVER_SUB}])
    textbox(s, ML, Inches(5.85), CW, Inches(0.70),
            [{"text": "The resume is the entry point. The ambition is the global AI career platform — connecting job seekers, employers, career services, education, and the talent ecosystem.",
              "size": 15, "color": COVER_SUB}])
    textbox(s, ML, Inches(6.70), CW, Inches(0.28),
            [{"text": "September 2026    |    New York  ·  Shanghai    |    China R&D  ·  U.S. commercialization  ·  Global distribution",
              "size": 12, "color": COVER_META}])


def s_exec(prs, n):
    s = blank(prs)
    header(s, "Company Overview", "51 Careers at a Glance",
           "Career services  →  AI Resume  →  AI Career Platform  →  Global Career Marketplace")
    round_rect(s, ML, Inches(1.88), Inches(6.20), Inches(4.10), fill=PALE)
    textbox(s, Inches(1.12), Inches(2.04), Inches(5.76), Inches(3.78), [
        {"text": "Who we are", "size": 14, "bold": True, "color": BLUE, "space_after": 6},
        {"text": "51 Careers is building a global AI Career Platform to help job seekers create better resumes, discover opportunities, improve job-search outcomes, and access professional career services.", "size": 14, "color": BODY, "space_after": 12},
        {"text": "Entry point", "size": 14, "bold": True, "color": BLUE, "space_after": 6},
        {"text": "An AI-powered Resume Platform that turns unstructured career experience into professional, quantified, job-specific resumes.", "size": 14, "color": BODY, "space_after": 12},
        {"text": "This round", "size": 14, "bold": True, "color": BLUE, "space_after": 6},
        {"text": "RMB 20 million (10% of RMB 200 million post-money) to accelerate product development, user acquisition, and global commercialization. Pre-money RMB 180 million.", "size": 14, "color": BODY},
    ])
    highlights = [
        "AI Resume demo already developed",
        "Initial testing completed with promising results",
        "Stronger professional positioning vs. general-purpose models",
        "Product can serve users globally from day one",
        "Existing career-services business: expertise + monetization",
        "U.S. market as the first international growth beachhead",
        "Multiple revenue streams beyond AI Resume subscriptions",
        "RMB 20M to fund product, acquisition, and global GTM",
    ]
    round_rect(s, Inches(7.30), Inches(1.88), Inches(5.13), Inches(4.10), fill=PALE)
    textbox(s, Inches(7.50), Inches(2.04), Inches(4.75), Inches(0.32),
            [{"text": "Investment highlights", "size": 14, "bold": True, "color": TITLE}])
    for i, t in enumerate(highlights):
        textbox(s, Inches(7.50), Inches(2.42) + i * Inches(0.42), Inches(4.75), Inches(0.40),
                [{"text": f"{i+1:02d}    {t}", "size": 13, "color": BODY}])
    stats = [
        ("2016", "Founded in New York"),
        ("RMB 20M", "Growth capital this round"),
        ("10%", "Equity at RMB 200M post"),
        ("RMB 180M", "Pre-money valuation"),
    ]
    for i, (big, lab) in enumerate(stats):
        x = ML + i * Inches(2.91)
        round_rect(s, x, Inches(6.10), Inches(2.80), Inches(0.84), fill=PALE)
        textbox(s, x + Inches(0.14), Inches(6.16), Inches(2.52), Inches(0.36),
                [{"text": big, "size": 16, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.14), Inches(6.50), Inches(2.52), Inches(0.34),
                [{"text": lab, "size": 12, "color": MUTED}])
    chrome(s, n)


def s_problem(prs, n):
    s = blank(prs)
    header(s, "The Problem", "Job Seeking Is Still Highly Inefficient",
           "For millions of candidates the process remains fragmented, expensive, and low-signal.")
    items = [
        ("Resume quality", "Most candidates struggle to communicate experience in a professional, results-oriented way."),
        ("Quantification", "People describe responsibilities rather than measurable achievements."),
        ("Personalization", "One generic resume rarely works across jobs and employers."),
        ("Job matching", "Candidates apply without a clear view of actual fit."),
        ("Career guidance", "Professional advice remains expensive and consultant-dependent."),
    ]
    for i, (t, b) in enumerate(items):
        x = ML + i * Inches(2.33)
        round_rect(s, x, Inches(1.90), Inches(2.21), Inches(4.10), fill=PALE)
        badge(s, x + Inches(0.18), Inches(2.10), i + 1)
        textbox(s, x + Inches(0.18), Inches(2.65), Inches(1.85), Inches(0.90),
                [{"text": t, "size": 16, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.18), Inches(3.60), Inches(1.85), Inches(2.15),
                [{"text": b, "size": 13.5, "color": BODY}])
    conclusion(s, "The result",
               "Low-quality applications  →  Low interview rates  →  Long job searches  →  High frustration")
    chrome(s, n)


def s_opportunity(prs, n):
    s = blank(prs)
    header(s, "The Opportunity", "AI Is Reshaping Career Services",
           "Generative AI has changed how people create, communicate, and consume professional information.")
    round_rect(s, ML, Inches(1.90), Inches(5.56), Inches(4.10), fill=PALE)
    textbox(s, Inches(1.12), Inches(2.10), Inches(5.20), Inches(0.36),
            [{"text": "Traditional model", "size": 18, "bold": True, "color": TITLE}])
    textbox(s, Inches(1.12), Inches(2.50), Inches(5.20), Inches(0.28),
            [{"text": "Human consultant", "size": 13, "bold": True, "color": BLUE}])
    for i, line in enumerate(["Manual analysis", "Manual resume writing", "Manual job search", "High cost", "Limited scalability"]):
        textbox(s, Inches(1.12), Inches(3.00) + i * Inches(0.56), Inches(5.20), Inches(0.52),
                [{"text": line, "size": 16, "color": BODY}])
    round_rect(s, Inches(6.87), Inches(1.90), Inches(5.56), Inches(4.10), fill=PALE)
    textbox(s, Inches(7.09), Inches(2.10), Inches(5.20), Inches(0.36),
            [{"text": "AI-enabled model", "size": 18, "bold": True, "color": TITLE}])
    textbox(s, Inches(7.09), Inches(2.50), Inches(5.20), Inches(0.28),
            [{"text": "Career platform", "size": 13, "bold": True, "color": BLUE}])
    for i, line in enumerate(["Automated career analysis", "AI resume generation", "Job matching", "Continuous optimization", "Global scalability"]):
        textbox(s, Inches(7.09), Inches(3.00) + i * Inches(0.56), Inches(5.20), Inches(0.52),
                [{"text": line, "size": 16, "color": BODY}])
    conclusion(s, "Our opportunity",
               "Transform high-cost, labor-intensive career services into a scalable AI-powered platform.")
    chrome(s, n)


def s_why_now(prs, n):
    s = blank(prs)
    header(s, "Why Now", "Three Structural Trends Are Converging",
           "AI adoption, data-driven hiring, and fragmented career services are moving at the same time.")
    trends = [
        ("Generative AI has reached mass adoption",
         "Consumers are increasingly comfortable using AI for professional and personal tasks. That adoption is the distribution opening for an AI career product."),
        ("Recruiting is becoming more data-driven",
         "Employers rely on structured information, keywords, measurable achievements, and automated screening. Candidates need a product that writes to that standard."),
        ("Career services remain fragmented",
         "The industry is still dependent on manual services and fragmented providers. That gap is the opening for a scalable AI + human platform."),
    ]
    for i, (t, b) in enumerate(trends):
        x = ML + i * Inches(3.91)
        round_rect(s, x, Inches(1.90), Inches(3.71), Inches(4.10), fill=PALE)
        badge(s, x + Inches(0.22), Inches(2.10), i + 1)
        textbox(s, x + Inches(0.22), Inches(2.68), Inches(3.27), Inches(1.15),
                [{"text": t, "size": 18, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.22), Inches(3.95), Inches(3.27), Inches(1.80),
                [{"text": b, "size": 14, "color": BODY}])
    conclusion(s, "51 Careers sits at the intersection", "AI   ×   Employment   ×   Career Services")
    chrome(s, n)


def s_tam(prs, n):
    s = blank(prs)
    header(s, "Market Size", "TAM  →  SAM  →  SOM  →  Three-Year Revenue",
           "From global job seekers to paid AI career users. Illustrative, for discussion.")
    cards = [
        ("TAM", "US$ 200B+", "RMB 1.35T+",
         "Global recruitment, talent-acquisition, and career-development spend."),
        ("SAM", "US$ 30B+", "RMB 200B+",
         "Digital / AI-enabled hiring tools, resume products, and online career preparation."),
        ("SOM", "US$ 1.5B+", "RMB 10B+",
         "AI resume and career-assistant demand in China, the U.S., and English-speaking Asia."),
    ]
    for i, (lab, usd, rmb, desc) in enumerate(cards):
        x = ML + i * Inches(3.91)
        round_rect(s, x, Inches(1.90), Inches(3.71), Inches(2.55), fill=PALE)
        textbox(s, x + Inches(0.22), Inches(2.04), Inches(3.27), Inches(0.26),
                [{"text": lab, "size": 13, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.22), Inches(2.32), Inches(3.27), Inches(0.50),
                [{"text": usd, "size": 26, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.22), Inches(2.84), Inches(3.27), Inches(0.26),
                [{"text": rmb, "size": 15, "bold": True, "color": BODY}])
        textbox(s, x + Inches(0.22), Inches(3.16), Inches(3.27), Inches(1.10),
                [{"text": desc, "size": 13.5, "color": MUTED}])
    round_rect(s, ML, Inches(4.58), CW, Inches(1.42), fill=PALE)
    textbox(s, Inches(1.12), Inches(4.68), CW - Inches(0.44), Inches(0.24),
            [{"text": "51 CAREERS  ·  2027E–2029E REVENUE", "size": 12, "bold": True, "color": BLUE}])
    kpis = [
        ("2027E", "US$ 2.0M", "≈ RMB 13.5M"),
        ("2028E", "US$ 8.0M", "≈ RMB 53.8M"),
        ("2029E", "US$ 28.0M", "≈ RMB 188.4M"),
    ]
    for i, (yr, a, b) in enumerate(kpis):
        x = Inches(1.12) + i * Inches(3.80)
        textbox(s, x, Inches(4.98), Inches(3.50), Inches(0.22),
                [{"text": yr, "size": 13, "bold": True, "color": BLUE}])
        textbox(s, x, Inches(5.20), Inches(3.50), Inches(0.42),
                [{"text": a, "size": 22, "bold": True, "color": TITLE}])
        textbox(s, x, Inches(5.62), Inches(3.50), Inches(0.26),
                [{"text": b, "size": 13, "color": MUTED}])
    conclusion(s, "Logic",
               "Global job seekers  →  addressable AI Resume users  →  paid conversion  →  stacked revenue (AI + services + commission + B2B + ads).")
    chrome(s, n)


def s_solution(prs, n):
    s = blank(prs)
    header(s, "Our Solution", "51 Careers AI Career Platform",
           "Not another AI writing tool — an AI-powered career infrastructure layer.")
    steps = [
        "Understand professional experience",
        "Build a stronger professional profile",
        "Generate optimized resumes",
        "Match with relevant opportunities",
        "Prepare for interviews",
        "Make better career decisions",
        "Access professional services",
    ]
    top, block_h, row_gap = 1.90, 5.00, 0.08
    row_h = (block_h - 6 * row_gap) / 7
    for i, step in enumerate(steps):
        y = top + i * (row_h + row_gap)
        round_rect(s, ML, Inches(y), Inches(7.15), Inches(row_h), fill=PALE)
        badge(s, Inches(1.08), Inches(y + (row_h - 0.38) / 2), i + 1)
        textbox(s, Inches(1.62), Inches(y), Inches(6.20), Inches(row_h),
                [{"text": step, "size": 16, "bold": True, "color": TITLE}], valign=MSO_ANCHOR.MIDDLE)
    round_rect(s, Inches(8.25), Inches(top), Inches(4.18), Inches(block_h), fill=PALE)
    textbox(s, Inches(8.47), Inches(2.08), Inches(3.76), Inches(0.26),
            [{"text": "Starting point", "size": 13, "bold": True, "color": BLUE}])
    textbox(s, Inches(8.47), Inches(2.42), Inches(3.76), Inches(0.70),
            [{"text": "AI Resume", "size": 28, "bold": True, "color": TITLE}])
    textbox(s, Inches(8.47), Inches(3.25), Inches(3.76), Inches(3.40),
            [{"text": "The resume is the first touchpoint between the user and the 51 Careers ecosystem — the on-ramp to AI Career tools, services, matching, and the marketplace.\n\nOne product. One profile. Multiple monetization paths across the career journey.",
              "size": 15, "color": BODY}])
    chrome(s, n)


def s_product(prs, n):
    s = blank(prs)
    header(s, "Product", "From Career Experience to Quantified Professional Value")
    round_rect(s, ML, Inches(1.55), Inches(3.70), Inches(5.40), fill=PALE)
    textbox(s, Inches(1.12), Inches(1.72), Inches(3.30), Inches(0.28),
            [{"text": "Users provide", "size": 13, "bold": True, "color": BLUE}])
    for i, item in enumerate(["Education", "Work experience", "Projects", "Skills", "Achievements", "Career goals", "Target job descriptions"]):
        textbox(s, Inches(1.12), Inches(2.12) + i * Inches(0.66), Inches(3.30), Inches(0.58),
                [{"text": item, "size": 16, "color": TITLE}])
    round_rect(s, Inches(4.80), Inches(1.55), Inches(3.70), Inches(5.40), fill=PALE)
    textbox(s, Inches(5.02), Inches(1.72), Inches(3.30), Inches(0.28),
            [{"text": "Our AI generates", "size": 13, "bold": True, "color": BLUE}])
    textbox(s, Inches(5.02), Inches(2.20), Inches(3.30), Inches(2.60), [
        {"text": "Professional", "size": 24, "bold": True, "color": TITLE, "space_after": 10},
        {"text": "+  Quantified", "size": 24, "bold": True, "color": TITLE, "space_after": 10},
        {"text": "+  Job-specific", "size": 24, "bold": True, "color": TITLE, "space_after": 10},
        {"text": "resume content", "size": 16, "color": MUTED},
    ])
    textbox(s, Inches(5.02), Inches(5.20), Inches(3.30), Inches(1.50),
            [{"text": "Optimized for ATS, role, employer, and geography — so one career profile can serve many applications.",
              "size": 14, "color": BODY}])
    caps = [
        "AI Resume Generation", "Resume Optimization", "Job-Specific Customization", "Achievement Quantification",
        "ATS Optimization", "Multiple Resume Versions", "Cover Letter Generation", "Career Profile Creation",
    ]
    for i, cap in enumerate(caps):
        y = Inches(1.55) + i * Inches(0.675)
        round_rect(s, Inches(8.70), y, Inches(3.73), Inches(0.62), fill=PALE)
        textbox(s, Inches(8.88), y, Inches(3.40), Inches(0.62),
                [{"text": cap, "size": 14, "bold": True, "color": TITLE}], valign=MSO_ANCHOR.MIDDLE)
    chrome(s, n)


def s_validation(prs, n):
    s = blank(prs)
    header(s, "Product Validation", "The Product Has Already Been Built and Tested")
    checks = [
        ("Demo developed", "A working AI Resume product is available for demonstration."),
        ("Initial testing completed", "Early users and internal reviews show promising quality."),
        ("Internal benchmark testing", "Compared against direct outputs from leading general-purpose AI models."),
    ]
    for i, (t, b) in enumerate(checks):
        x = ML + i * Inches(3.91)
        round_rect(s, x, Inches(1.55), Inches(3.71), Inches(1.95), fill=PALE)
        textbox(s, x + Inches(0.22), Inches(1.70), Inches(3.27), Inches(0.40),
                [{"text": t, "size": 16, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.22), Inches(2.18), Inches(3.27), Inches(1.10),
                [{"text": b, "size": 14, "color": BODY}])
    textbox(s, ML, Inches(3.64), CW, Inches(0.28),
            [{"text": "Initial testing indicates stronger performance in", "size": 13, "bold": True, "color": BLUE}])
    dims = ["Professional positioning", "Achievement quantification", "Job relevance",
            "Career-experience extraction", "Resume structure"]
    for i, d in enumerate(dims):
        x = ML + i * Inches(2.33)
        round_rect(s, x, Inches(4.00), Inches(2.21), Inches(1.95), fill=PALE)
        textbox(s, x + Inches(0.14), Inches(4.45), Inches(1.93), Inches(1.15),
                [{"text": d, "size": 15, "bold": True, "color": TITLE}])
    conclusion(s, "Important note",
               "Current results are internal. Next: larger-scale, structured, blind A/B tests versus general-purpose models.")
    chrome(s, n)


def s_tech(prs, n):
    s = blank(prs)
    header(s, "Technology", "We Are Not Simply Building an LLM Wrapper")
    round_rect(s, ML, Inches(1.55), CW, Inches(0.78), fill=PALE)
    textbox(s, Inches(1.12), Inches(1.70), CW - Inches(0.44), Inches(0.50),
            [{"text": "LLM  +  Career Knowledge  +  Resume Framework  +  Job Intelligence  +  User Data",
              "size": 16, "bold": True, "color": TITLE}])
    chips = [
        "Career-experience extraction", "Achievement quantification", "Job-description analysis", "Resume optimization",
        "Industry terminology", "Role-specific requirements", "Geographic resume conventions", "ATS optimization",
    ]
    for i, c in enumerate(chips):
        col, row = i % 4, i // 4
        x = ML + col * Inches(2.91)
        y = Inches(2.50) + row * Inches(1.05)
        round_rect(s, x, y, Inches(2.80), Inches(0.92), fill=PALE)
        textbox(s, x + Inches(0.16), y + Inches(0.20), Inches(2.48), Inches(0.54),
                [{"text": c, "size": 14, "bold": True, "color": TITLE}], valign=MSO_ANCHOR.MIDDLE)
    loop = ["User input", "AI analysis", "Resume gen.", "Application", "Feedback", "Model upgrade", "Better outcomes", "More users"]
    textbox(s, ML, Inches(4.70), CW, Inches(0.26),
            [{"text": "Data feedback loop", "size": 13, "bold": True, "color": BLUE}])
    for i, step in enumerate(loop):
        x = ML + i * Inches(1.45)
        round_rect(s, x, Inches(5.02), Inches(1.35), Inches(0.98), fill=PALE)
        textbox(s, x + Inches(0.08), Inches(5.10), Inches(1.20), Inches(0.26),
                [{"text": f"{i+1:02d}", "size": 12, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.08), Inches(5.38), Inches(1.20), Inches(0.52),
                [{"text": step, "size": 12, "bold": True, "color": TITLE}])
    conclusion(s, "Compounding advantage",
               "Each cycle builds proprietary career intelligence that generic chat models do not accumulate.")
    chrome(s, n)


def s_global_market(prs, n):
    s = blank(prs)
    header(s, "Global Market", "AI Resume Is a Naturally Global Product",
           "A resume is a universal component of the global employment market.")
    markets = ["United States", "Canada", "Central America", "United Kingdom", "Australia",
               "Singapore", "Hong Kong", "China", "Europe", "Middle East"]
    for i, m in enumerate(markets):
        col, row = i % 5, i // 5
        x = ML + col * Inches(2.33)
        y = Inches(1.90) + row * Inches(2.00)
        round_rect(s, x, y, Inches(2.21), Inches(1.85), fill=PALE)
        textbox(s, x + Inches(0.12), y + Inches(0.60), Inches(1.97), Inches(0.70),
                [{"text": m, "size": 16, "bold": True, "color": TITLE}])
    conclusion(s, "The universal question",
               "“How do I present my professional value to employers?” — a global problem with a scalable product.")
    chrome(s, n)


def s_strategy(prs, n):
    s = blank(prs)
    header(s, "Product Strategy", "One Platform. Multiple Labor Markets.")
    dims = [
        ("Geography",
         "U.S.  ·  Canada  ·  Central America  ·  U.K.  ·  Australia  ·  Europe  ·  Asia  ·  Middle East"),
        ("Industry",
         "Applicable across industries — including Technology, Fintech, Financial Services, Advanced Manufacturing, Consulting, Healthcare, Marketing, Education, and more."),
        ("Career level", "Student / Graduate / Professional / Manager / Executive"),
        ("Target role", "Software Engineer / PM / Data Scientist / Marketing / Finance / more"),
        ("Target employer", "Startup / SME / Enterprise / Multinational"),
    ]
    for i, (t, b) in enumerate(dims):
        y = Inches(1.55) + i * Inches(1.06)
        round_rect(s, ML, y, CW, Inches(0.96), fill=PALE)
        textbox(s, Inches(1.12), y + Inches(0.12), Inches(2.70), Inches(0.72),
                [{"text": t, "size": 18, "bold": True, "color": TITLE}], valign=MSO_ANCHOR.MIDDLE)
        textbox(s, Inches(4.00), y + Inches(0.12), Inches(8.10), Inches(0.72),
                [{"text": b, "size": 15, "color": BODY}], valign=MSO_ANCHOR.MIDDLE)
    chrome(s, n)


def s_phases(prs, n):
    s = blank(prs)
    header(s, "Product Roadmap", "Resume Is the Entry Point — Not the Destination")
    phases = [
        ("Phase 1", "AI Resume", "“Help me create a better resume.”"),
        ("Phase 2", "AI Career Assistant", "“What jobs should I apply for?”"),
        ("Phase 3", "AI Career Agent", "“Help me manage my entire job search.”"),
        ("Phase 4", "Global Platform", "Job seekers × employers × services × education × talent ecosystem"),
    ]
    for i, (p, t, b) in enumerate(phases):
        x = ML + i * Inches(2.91)
        round_rect(s, x, Inches(1.55), Inches(2.80), Inches(3.40), fill=PALE)
        badge(s, x + Inches(0.18), Inches(1.72), i + 1)
        textbox(s, x + Inches(0.18), Inches(2.28), Inches(2.44), Inches(0.28),
                [{"text": p, "size": 13, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.18), Inches(2.62), Inches(2.44), Inches(0.80),
                [{"text": t, "size": 18, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.18), Inches(3.50), Inches(2.44), Inches(1.20),
                [{"text": b, "size": 14, "color": BODY}])
    caps = ["Resume creation", "Job discovery", "Job matching", "Application optimization",
            "Cover letters", "Interview prep", "Career planning"]
    for i, c in enumerate(caps):
        x = ML + i * Inches(1.66)
        round_rect(s, x, Inches(5.10), Inches(1.55), Inches(1.70), fill=PALE)
        textbox(s, x + Inches(0.08), Inches(5.45), Inches(1.39), Inches(1.10),
                [{"text": c, "size": 13, "bold": True, "color": TITLE}])
    chrome(s, n)


def s_model(prs, n):
    s = blank(prs)
    header(s, "Business Model", "Multiple Monetization Layers Across the Career Journey")
    streams = [
        ("AI product", "Subscription, premium AI features, resume packages, AI Career tools."),
        ("Career services", "Resume consulting, career consulting, interview coaching, premium services."),
        ("Referral & commission", "CPA / CPS / revenue share with education, training, and recruitment partners."),
        ("B2B / employers", "Talent acquisition, employer branding, recruitment campaigns, job promotion."),
        ("Advertising", "As traffic scales: relevant brands → targeted career audience."),
        ("Sponsorship", "Career fairs, conferences, recruitment events, AI career events."),
    ]
    for i, (t, b) in enumerate(streams):
        col, row = i % 3, i // 3
        x = ML + col * Inches(3.91)
        y = Inches(1.55) + row * Inches(2.65)
        round_rect(s, x, y, Inches(3.71), Inches(2.50), fill=PALE)
        badge(s, x + Inches(0.20), y + Inches(0.22), i + 1)
        textbox(s, x + Inches(0.72), y + Inches(0.22), Inches(2.78), Inches(0.40),
                [{"text": t, "size": 16, "bold": True, "color": TITLE}], valign=MSO_ANCHOR.MIDDLE)
        textbox(s, x + Inches(0.20), y + Inches(0.85), Inches(3.30), Inches(1.40),
                [{"text": b, "size": 15, "color": BODY}])
    chrome(s, n)


def s_flywheel(prs, n):
    s = blank(prs)
    header(s, "The Flywheel", "One User Can Generate Multiple Revenue Opportunities")
    steps = ["Free AI Resume", "User acquisition", "Career profile", "AI Career tools", "Premium subscription",
             "Career services", "Job matching", "Recruitment / referral", "Advertising", "Employer services"]
    for i, st in enumerate(steps):
        col, row = i % 5, i // 5
        x = ML + col * Inches(2.33)
        y = Inches(1.55) + row * Inches(2.20)
        round_rect(s, x, y, Inches(2.21), Inches(2.05), fill=PALE)
        textbox(s, x + Inches(0.14), y + Inches(0.18), Inches(1.93), Inches(0.32),
                [{"text": f"{i+1:02d}", "size": 14, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.14), y + Inches(0.60), Inches(1.93), Inches(1.20),
                [{"text": st, "size": 16, "bold": True, "color": TITLE}])
    conclusion(s, "Therefore",
               "Revenue is not limited to the first AI Resume transaction. The objective is lifetime value per career user.")
    chrome(s, n)


def s_acquisition(prs, n):
    s = blank(prs)
    header(s, "User Acquisition", "Digital  +  Offline  +  Organic")
    cols = [
        ("Digital", ["Google", "Meta", "TikTok", "YouTube", "LinkedIn", "Search / social", "Content marketing"]),
        ("Organic", ["SEO", "AI-search optimization", "Career content", "Educational content", "Referral programs", "Affiliate partnerships"]),
        ("Offline", ["Universities", "Career fairs", "Student organizations", "Professional orgs", "Community events", "Employer events"]),
    ]
    for i, (t, items) in enumerate(cols):
        x = ML + i * Inches(3.91)
        round_rect(s, x, Inches(1.55), Inches(3.71), Inches(4.45), fill=PALE)
        textbox(s, x + Inches(0.22), Inches(1.72), Inches(3.27), Inches(0.36),
                [{"text": t, "size": 18, "bold": True, "color": TITLE}])
        for j, it in enumerate(items):
            textbox(s, x + Inches(0.22), Inches(2.20) + j * Inches(0.52), Inches(3.27), Inches(0.48),
                    [{"text": it, "size": 15, "color": BODY}])
    conclusion(s, "Core growth strategy",
               "Free AI product  →  Low-friction acquisition  →  Registration  →  Engagement  →  Monetization")
    chrome(s, n)


def s_geo(prs, n):
    s = blank(prs)
    header(s, "Geography", "China as R&D Base  ·  U.S. as Commercialization Market")
    geos = [
        ("1", "China  ·  R&D and product",
         "Product, design, engineering, AI, testing, marketing operations, B2B business development."),
        ("2", "United States  ·  International GTM",
         "User acquisition, local partnerships, employer and university relationships, offline events, brand building."),
        ("3", "Global expansion",
         "Once product-market fit is established: North America → UK → Australia → Europe → Asia → other markets."),
    ]
    for i, (num, t, b) in enumerate(geos):
        y = Inches(1.55) + i * Inches(1.78)
        round_rect(s, ML, y, CW, Inches(1.66), fill=PALE)
        badge(s, Inches(1.12), y + Inches(0.22), num)
        textbox(s, Inches(1.70), y + Inches(0.18), Inches(10.4), Inches(0.42),
                [{"text": t, "size": 18, "bold": True, "color": TITLE}])
        textbox(s, Inches(1.70), y + Inches(0.68), Inches(10.4), Inches(0.78),
                [{"text": b, "size": 15, "color": BODY}])
    chrome(s, n)


def s_compete(prs, n):
    s = blank(prs)
    header(s, "Competitive Landscape", "Career Infrastructure — Not Another Resume Builder")
    quads = [
        ("Job platforms", "Large user base  ·  Limited personalized career intelligence"),
        ("General-purpose AI", "High AI capability  ·  Limited career specialization"),
        ("Traditional career services", "High human service  ·  Low scalability"),
        ("AI resume tools", "Resume-focused  ·  Limited ecosystem"),
    ]
    for i, (t, b) in enumerate(quads):
        col, row = i % 2, i // 2
        x = ML + col * Inches(5.87)
        y = Inches(1.55) + row * Inches(2.20)
        round_rect(s, x, y, Inches(5.66), Inches(2.05), fill=PALE)
        textbox(s, x + Inches(0.24), y + Inches(0.28), Inches(5.18), Inches(0.42),
                [{"text": t, "size": 18, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.24), y + Inches(0.82), Inches(5.18), Inches(0.95),
                [{"text": b, "size": 16, "color": BODY}])
    conclusion(s, "51 Careers",
               "AI + career expertise + services + marketplace. We are building career infrastructure, not another resume builder.")
    chrome(s, n)


def s_advantages(prs, n):
    s = blank(prs)
    header(s, "Competitive Position", "Why Customers Will Choose 51 Careers")
    items = [
        ("Existing career-services foundation", "We understand the industry from an operating perspective."),
        ("AI product already in development", "Not a concept or a slide-deck company. Demo built and tested."),
        ("Global scalability", "AI Resume can be distributed worldwide through digital channels."),
        ("Multiple monetization paths", "Subscription + services + commission + B2B + ads + sponsorship."),
        ("AI + human hybrid", "AI for scalable work; humans for high-value, complex services."),
        ("Platform potential", "The long-term opportunity extends beyond resumes into the career ecosystem."),
    ]
    for i, (t, b) in enumerate(items):
        col, row = i % 3, i // 3
        x = ML + col * Inches(3.91)
        y = Inches(1.55) + row * Inches(2.65)
        round_rect(s, x, y, Inches(3.71), Inches(2.50), fill=PALE)
        badge(s, x + Inches(0.20), y + Inches(0.22), i + 1)
        textbox(s, x + Inches(0.20), y + Inches(0.78), Inches(3.30), Inches(0.70),
                [{"text": t, "size": 16, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.20), y + Inches(1.52), Inches(3.30), Inches(0.78),
                [{"text": b, "size": 14, "color": BODY}])
    chrome(s, n)


def s_leadership(prs, n):
    s = blank(prs)
    header(s, "Leadership", "Executive Leadership Team",
           "Founder-led. Globally distributed.")
    leaders = [
        ("Rocky Chen", "Founder & CEO",
         "Founded 51 Careers in New York in 2016 and built it from a boutique consultancy into a multi-market career-services and technology business. Founder of 51 Careers.AI; investor and Managing Director at publicly listed Helio. Architect of the AI-first strategy, including the 2025 Alibaba Cloud AI partnership."),
        ("Stephanie Li", "Co-Founder & CFO",
         "Co-founder since 2016. U.S. CPA; M.S. Accounting, Pace University. Oversees financial strategy, corporate governance, and sustainable growth for 51 Careers and 51 Careers.AI."),
        ("Gavin Ding", "Co-Founder, CTO & COO",
         "Serial entrepreneur (10+ years) across SaaS, AI, and digital business; B.S. Computer Science, ECUST. Prior CTO roles and multiple co-founded ventures. Leads technology, operations, and AI."),
        ("Robin Zhu", "Head of Product",
         "Leads product strategy, architecture, and development of the AI-powered career platform. Former CTO & Director of Product at Ci Finance; senior product roles at CPIC, Allinpay, and Noah."),
        ("Chris Lin", "North American Partner",
         "Full-stack engineer; previously at Amazon Web Services (AWS). B.S. Northwestern; M.S. Robotics. Extensive hiring-panel experience and professional mentorship."),
        ("Jon Serbin", "Senior Advisor",
         "Harvard and MIT; former senior executive at Morgan Stanley; founder of Cedar. 40+ years in technology M&A and capital raising. Advises 51 Careers on growth and expansion."),
    ]
    for i, (name, title, bio) in enumerate(leaders):
        col, row = i % 3, i // 3
        x = ML + col * Inches(3.91)
        y = Inches(1.88) + row * Inches(2.52)
        round_rect(s, x, y, Inches(3.71), Inches(2.40), fill=PALE)
        textbox(s, x + Inches(0.18), y + Inches(0.12), Inches(3.35), Inches(0.34),
                [{"text": name, "size": 16, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.18), y + Inches(0.44), Inches(3.35), Inches(0.24),
                [{"text": title.upper(), "size": 11, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.18), y + Inches(0.72), Inches(3.35), Inches(1.54),
                [{"text": bio, "size": 12, "color": BODY}])
    chrome(s, n)


def s_team(prs, n):
    s = blank(prs)
    header(s, "Team", "Core Founding Team in Place  ·  Planned Hiring Funded by This Round")
    round_rect(s, ML, Inches(1.55), Inches(3.70), Inches(5.40), fill=PALE)
    textbox(s, Inches(1.12), Inches(1.72), Inches(3.30), Inches(0.24),
            [{"text": "Existing", "size": 13, "bold": True, "color": BLUE}])
    textbox(s, Inches(1.12), Inches(2.00), Inches(3.30), Inches(0.36),
            [{"text": "Founding team", "size": 20, "bold": True, "color": TITLE}])
    textbox(s, Inches(1.12), Inches(2.42), Inches(3.30), Inches(0.70),
            [{"text": "4", "size": 44, "bold": True, "color": BLUE}])
    for i, line in enumerate([
        "In place today",
        "Product, capital, and go-to-market leadership",
        "Additional roles are not yet hired",
        "This round funds the planned operating team",
    ]):
        textbox(s, Inches(1.12), Inches(3.25) + i * Inches(0.82), Inches(3.30), Inches(0.74),
                [{"text": line, "size": 14, "color": BODY}])
    planned = [
        ("China technology", "7",
         [("Product Manager", "1"), ("Designer", "1"), ("QA / Testing", "1"),
          ("Front-End Engineer", "1"), ("Back-End Engineers", "2"), ("AI Engineer", "1")]),
        ("China growth", "2+",
         [("Marketing / Performance", "1"), ("B2B Business Development", "1+")]),
        ("U.S. growth", "3",
         [("Performance Marketing", "1"), ("Local / Field Marketing", "2")]),
    ]
    for i, (title, count, roles) in enumerate(planned):
        x = Inches(4.80) + i * Inches(2.75)
        round_rect(s, x, Inches(1.55), Inches(2.60), Inches(5.40), fill=PALE)
        textbox(s, x + Inches(0.16), Inches(1.70), Inches(2.28), Inches(0.22),
                [{"text": "Planned", "size": 12, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.16), Inches(1.96), Inches(2.28), Inches(0.50),
                [{"text": title, "size": 15, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.16), Inches(2.48), Inches(2.28), Inches(0.42),
                [{"text": count, "size": 26, "bold": True, "color": BLUE}])
        for j, (role, nhead) in enumerate(roles):
            textbox(s, x + Inches(0.16), Inches(3.05) + j * Inches(0.58), Inches(1.70), Inches(0.52),
                    [{"text": role, "size": 12, "color": BODY}])
            textbox(s, x + Inches(1.80), Inches(3.05) + j * Inches(0.58), Inches(0.62), Inches(0.52),
                    [{"text": nhead, "size": 12, "bold": True, "color": BLUE, "align": PP_ALIGN.RIGHT}])
    chrome(s, n)


def s_funds(prs, n):
    s = blank(prs)
    header(s, "Use of Funds", "RMB 20 Million Growth Financing",
           "RMB 13.1M is a planned personnel + advertising budget. Remaining capital funds infrastructure, AI, operations, events, expansion and working capital.")
    alloc = [
        ("Marketing & user acquisition", "40%", "RMB 8.0M"),
        ("Product & AI development", "22.5%", "RMB 4.5M"),
        ("U.S. market expansion", "12.5%", "RMB 2.5M"),
        ("China market & B2B", "7.5%", "RMB 1.5M"),
        ("Brand & events", "7.5%", "RMB 1.5M"),
        ("Operations / legal / finance", "5%", "RMB 1.0M"),
        ("Cash reserve", "5%", "RMB 1.0M"),
    ]
    data = CategoryChartData()
    data.categories = [a[0] for a in alloc]
    data.add_series("Use of Funds", (40, 22.5, 12.5, 7.5, 7.5, 5, 5))
    chart = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(0.40), Inches(1.88), Inches(5.20), Inches(2.95), data).chart
    chart.has_legend = False
    for i, (name, pct, amt) in enumerate(alloc):
        y = Inches(1.90) + i * Inches(0.42)
        textbox(s, Inches(5.60), y, Inches(4.80), Inches(0.40),
                [{"text": name, "size": 14, "color": BODY}])
        textbox(s, Inches(10.35), y, Inches(2.08), Inches(0.40),
                [{"text": f"{pct}   ·   {amt}", "size": 14, "bold": True, "color": TITLE, "align": PP_ALIGN.RIGHT}])
    kpis = [
        ("Personnel", "RMB 5.1M", "Planned hiring — not current payroll"),
        ("China ads", "RMB 1.0M", "Planned annual spend"),
        ("U.S. ads", "RMB 7.0M", "Planned annual spend"),
        ("P+A budget", "RMB 13.1M", "Planned personnel + advertising"),
    ]
    for i, (lab, big, sub) in enumerate(kpis):
        x = ML + i * Inches(2.91)
        round_rect(s, x, Inches(5.05), Inches(2.80), Inches(1.90), fill=PALE)
        textbox(s, x + Inches(0.16), Inches(5.16), Inches(2.48), Inches(0.24),
                [{"text": lab.upper(), "size": 12, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.16), Inches(5.42), Inches(2.48), Inches(0.42),
                [{"text": big, "size": 20, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.16), Inches(5.88), Inches(2.48), Inches(0.90),
                [{"text": sub, "size": 13, "color": MUTED}])
    chrome(s, n)


def s_roadmap(prs, n):
    s = blank(prs)
    header(s, "Growth Roadmap", "Validation  →  Growth  →  Global Expansion")
    phases = [
        ("Phase 1  ·  0–6 months", "Product validation",
         ["Launch AI Resume", "Improve AI quality", "Structured A/B testing", "Initial CAC benchmarks", "Conversion funnel", "Validate pricing"]),
        ("Phase 2  ·  6–12 months", "Growth",
         ["Scale paid acquisition", "Build SEO engine", "Expand U.S. acquisition", "B2B partnerships", "More AI Career features", "Expand monetization"]),
        ("Phase 3  ·  12–24 months", "Global expansion",
         ["More English-speaking markets", "Localized products", "Global partnerships", "Employer ecosystem", "Marketplace capabilities", "Central America expansion"]),
    ]
    for i, (lab, title, items) in enumerate(phases):
        x = ML + i * Inches(3.91)
        round_rect(s, x, Inches(1.55), Inches(3.71), Inches(5.40), fill=PALE)
        badge(s, x + Inches(0.20), Inches(1.72), i + 1)
        textbox(s, x + Inches(0.72), Inches(1.72), Inches(2.78), Inches(0.40),
                [{"text": lab, "size": 13, "bold": True, "color": BLUE}], valign=MSO_ANCHOR.MIDDLE)
        textbox(s, x + Inches(0.20), Inches(2.24), Inches(3.31), Inches(0.42),
                [{"text": title, "size": 18, "bold": True, "color": TITLE}])
        for j, it in enumerate(items):
            textbox(s, x + Inches(0.20), Inches(2.78) + j * Inches(0.64), Inches(3.31), Inches(0.58),
                    [{"text": it, "size": 14, "color": BODY}])
    chrome(s, n)


def s_model_3yr(prs, n):
    s = blank(prs)
    header(s, "Financial Model", "Users, Revenue, and Profitability",
           "Illustrative 2027E–2029E operating plan. Figures in US$.")
    panels = [
        ("Users",
         [("Users", "1.0M", "3.0M", "7.0M"),
          ("Registered users", "300K", "1.0M", "2.5M"),
          ("Monthly active users", "150K", "500K", "1.25M"),
          ("Paid users", "15K", "60K", "180K"),
          ("Paid conversion", "5.0%", "6.0%", "7.2%")]),
        ("Revenue",
         [("AI revenue", "$0.70M", "$2.40M", "$8.40M"),
          ("Career services", "$0.50M", "$1.60M", "$5.60M"),
          ("Commission", "$0.30M", "$1.20M", "$4.20M"),
          ("Advertising", "$0.20M", "$0.80M", "$2.80M"),
          ("B2B revenue", "$0.30M", "$2.00M", "$7.00M"),
          ("Total revenue", "$2.00M", "$8.00M", "$28.00M")]),
        ("Profitability",
         [("Gross margin", "65%", "70%", "75%"),
          ("Gross profit", "$1.30M", "$5.60M", "$21.00M"),
          ("EBITDA margin", "–50%", "–15%", "15%"),
          ("EBITDA", "–$1.00M", "–$1.20M", "$4.20M"),
          ("Profitability", "Loss", "Loss", "Profitable")]),
    ]
    years = ["2027E", "2028E", "2029E"]
    for i, (title, rows) in enumerate(panels):
        x = ML + i * Inches(3.91)
        round_rect(s, x, Inches(1.88), Inches(3.71), Inches(5.07), fill=PALE)
        textbox(s, x + Inches(0.16), Inches(2.00), Inches(3.39), Inches(0.28),
                [{"text": title, "size": 16, "bold": True, "color": TITLE}])
        for yi, yr in enumerate(years):
            textbox(s, x + Inches(1.10) + yi * Inches(0.85), Inches(2.32), Inches(0.85), Inches(0.24),
                    [{"text": yr, "size": 11, "bold": True, "color": MUTED, "align": PP_ALIGN.CENTER}])
        row_h = 4.20 / max(len(rows), 1)
        for ri, (metric, *vals) in enumerate(rows):
            y = Inches(2.62) + ri * Inches(row_h)
            strong = metric in {"Total revenue", "Profitability", "Users"}
            textbox(s, x + Inches(0.16), y, Inches(1.00), Inches(row_h - 0.04),
                    [{"text": metric, "size": 11, "bold": strong, "color": MUTED}])
            for vi, val in enumerate(vals):
                accent = BLUE if (metric == "Profitability" and val == "Profitable") else TITLE
                textbox(s, x + Inches(1.10) + vi * Inches(0.85), y, Inches(0.85), Inches(row_h - 0.04),
                        [{"text": val, "size": 13, "bold": True, "color": accent, "align": PP_ALIGN.CENTER}])
    chrome(s, n)


def s_investment(prs, n):
    s = blank(prs)
    header(s, "The Offering", "The Next Generation of Career Services Will Be AI + Human + Platform")
    axes = [
        ("Generative AI", "Professional content creation is now cheap, fast, and widely adopted."),
        ("Global employment", "Talent is mobile. English-language careers travel across markets."),
        ("Career services", "Advice is still manual, expensive, and hard to scale."),
        ("Digital user acquisition", "A free AI Resume is a low-friction, global top-of-funnel."),
    ]
    for i, (a, b) in enumerate(axes):
        col, row = i % 2, i // 2
        x = ML + col * Inches(5.87)
        y = Inches(1.55) + row * Inches(2.20)
        round_rect(s, x, y, Inches(5.66), Inches(2.05), fill=PALE)
        textbox(s, x + Inches(0.24), y + Inches(0.28), Inches(5.18), Inches(0.46),
                [{"text": a, "size": 20, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.24), y + Inches(0.86), Inches(5.18), Inches(0.95),
                [{"text": b, "size": 16, "color": BODY}])
    conclusion(s, "The opportunity",
               "The AI Resume is the entry point; the ambition is the global AI career platform.")
    chrome(s, n)


def s_financing(prs, n):
    s = blank(prs)
    header(s, "Financing", "Current Financing Round")
    terms = [
        ("Raise", "RMB 20 million", "≈ US$2.97 million"),
        ("Pre-money", "RMB 180 million", "≈ US$26.7 million"),
        ("Post-money", "RMB 200 million", "≈ US$29.7 million"),
        ("Equity offered", "10%", "Primary preferred"),
        ("Structure", "Preferred equity", "Indicative; subject to docs"),
        ("Use of capital", "Product, users, global GTM", "Monetization + ecosystem"),
    ]
    for i, (lab, a, b) in enumerate(terms):
        col, row = i % 3, i // 3
        x = ML + col * Inches(3.91)
        y = Inches(1.55) + row * Inches(2.20)
        round_rect(s, x, y, Inches(3.71), Inches(2.05), fill=PALE)
        textbox(s, x + Inches(0.22), y + Inches(0.20), Inches(3.27), Inches(0.28),
                [{"text": lab.upper(), "size": 13, "bold": True, "color": BLUE}])
        textbox(s, x + Inches(0.22), y + Inches(0.54), Inches(3.27), Inches(0.70),
                [{"text": a, "size": 22, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.22), y + Inches(1.28), Inches(3.27), Inches(0.55),
                [{"text": b, "size": 15, "color": MUTED}])
    conclusion(s, "Use of proceeds",
               "Build the product, acquire users, expand globally, build monetization, and develop the career ecosystem.")
    chrome(s, n)


def s_what_20m(prs, n):
    s = blank(prs)
    header(s, "Investor Outcomes", "12–18 Month Targets This Round Is Designed to Fund")
    outcomes = [
        ("Public launch", "AI Resume live for users in China and the U.S."),
        ("50–100k users", "Registered users; funnel instrumented end-to-end."),
        ("Blind A/B tests", "Statistically meaningful quality vs. general-purpose models."),
        ("Unit economics", "CAC, conversion, ARPU, and payback benchmarks."),
        ("U.S. acquisition", "Paid digital + university/offline channels operating."),
        ("First paid cohort", "AI subscriptions and career-services attach live."),
        ("B2B pipeline", "Employer and education-partner conversations in market."),
        ("Next-round pack", "Traction, model, and data room ready for the following raise."),
    ]
    for i, (t, b) in enumerate(outcomes):
        col, row = i % 4, i // 4
        x = ML + col * Inches(2.91)
        y = Inches(1.55) + row * Inches(2.20)
        round_rect(s, x, y, Inches(2.80), Inches(2.05), fill=PALE)
        badge(s, x + Inches(0.16), y + Inches(0.16), i + 1)
        textbox(s, x + Inches(0.16), y + Inches(0.64), Inches(2.48), Inches(0.48),
                [{"text": t, "size": 15, "bold": True, "color": TITLE}])
        textbox(s, x + Inches(0.16), y + Inches(1.16), Inches(2.48), Inches(0.72),
                [{"text": b, "size": 13, "color": BODY}])
    conclusion(s, "RMB 20M",
               "Product + users + benchmarks  →  revenue  →  next round  →  valuation step-up.")
    chrome(s, n)


def s_vision(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(COVER_BG, Inches(0), Inches(0), SW, SH)
    logo(s, light=True, cover=True)
    textbox(s, ML, Inches(1.70), Inches(11.20), Inches(0.28),
            [{"text": "VISION", "size": 13, "bold": True, "color": BLUE_LT}])
    textbox(s, ML, Inches(2.05), Inches(11.20), Inches(0.80),
            [{"text": "We are not building another AI resume tool.",
              "size": 26, "bold": True, "color": WHITE}])
    textbox(s, ML, Inches(2.90), Inches(11.20), Inches(0.80),
            [{"text": "We are building the Global AI Career Platform.",
              "size": 26, "bold": True, "color": BLUE_LT}])
    pillars = [
        ("China R&D", "Product, design, engineering, and AI built from the China base."),
        ("U.S. commercialization", "The first international growth market for acquisition and brand."),
        ("Global distribution", "A resume product that can serve English-speaking labor markets from day one."),
    ]
    for i, (t, b) in enumerate(pillars):
        x = ML + i * Inches(3.91)
        outlined_rect(s, x, Inches(3.90), Inches(3.71), Inches(2.15), COVER_CARD, COVER_LINE, 0.75)
        textbox(s, x + Inches(0.22), Inches(4.08), Inches(3.27), Inches(0.42),
                [{"text": t, "size": 16, "bold": True, "color": BLUE_LT}])
        textbox(s, x + Inches(0.22), Inches(4.55), Inches(3.27), Inches(1.25),
                [{"text": b, "size": 14, "color": COVER_SUB}])
    textbox(s, ML, Inches(6.25), CW, Inches(0.28),
            [{"text": "AI   ×   CAREER   ×   TALENT   ×   OPPORTUNITY", "size": 14, "color": BLUE_SOFT}])
    textbox(s, ML, Inches(6.70), Inches(8.50), Inches(0.26),
            [{"text": "51 Careers.AI   |   Proprietary and Confidential", "size": 10, "color": COVER_META}])
    textbox(s, Inches(12.10), Inches(6.70), Inches(0.35), Inches(0.26),
            [{"text": str(n), "size": 10, "color": COVER_META, "align": PP_ALIGN.RIGHT}])


def s_contact(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(COVER_BG, Inches(0), Inches(0), SW, SH)
    logo(s, light=True, cover=True)
    textbox(s, ML, Inches(1.70), Inches(11.20), Inches(0.28),
            [{"text": "THANK YOU", "size": 13, "bold": True, "color": BLUE_LT}])
    textbox(s, ML, Inches(2.05), Inches(11.20), Inches(0.70),
            [{"text": "Questions and Discussion", "size": 32, "bold": True, "color": WHITE}])
    textbox(s, ML, Inches(2.80), Inches(11.20), Inches(0.40),
            [{"text": "“Our mission is to make great jobs accessible to everyone.”",
              "size": 16, "color": COVER_SUB}])
    cards = [
        ("EMAIL", "rocky@helio.space"),
        ("PHONE", "19921169641"),
        ("WEB", "www.helio.space"),
    ]
    for i, (lab, val) in enumerate(cards):
        x = ML + i * Inches(3.91)
        outlined_rect(s, x, Inches(3.40), Inches(3.71), Inches(1.85), COVER_CARD, COVER_LINE, 0.75)
        textbox(s, x + Inches(0.22), Inches(3.58), Inches(3.27), Inches(0.28),
                [{"text": lab, "size": 12, "bold": True, "color": BLUE_LT}])
        textbox(s, x + Inches(0.22), Inches(4.00), Inches(3.27), Inches(0.95),
                [{"text": val, "size": 16, "bold": True, "color": WHITE}])
    outlined_rect(s, ML, Inches(5.45), CW, Inches(1.20), COVER_CARD, COVER_LINE, 0.75)
    textbox(s, Inches(1.12), Inches(5.60), CW - Inches(0.44), Inches(0.32),
            [{"text": "Rocky Chen  ·  Founder & Chief Executive Officer",
              "size": 16, "bold": True, "color": WHITE}])
    textbox(s, Inches(1.12), Inches(6.00), CW - Inches(0.44), Inches(0.42),
            [{"text": "info@helio.space   ·   (510) 545-2666   ·   2448 Sixth Street, Berkeley, CA 94710",
              "size": 14, "color": COVER_SUB}])
    textbox(s, Inches(10.20), Inches(7.08), Inches(2.25), Inches(0.26),
            [{"text": "Proprietary and Confidential", "size": 8.5, "color": COVER_META, "align": PP_ALIGN.RIGHT}])


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    s_cover(prs)
    builders = [
        s_exec, s_problem, s_opportunity, s_why_now, s_tam, s_solution, s_product,
        s_validation, s_tech, s_global_market, s_strategy, s_phases, s_model,
        s_flywheel, s_acquisition, s_geo, s_compete, s_advantages, s_leadership, s_team,
        s_funds, s_roadmap, s_model_3yr, s_investment, s_financing,
        s_what_20m, s_vision, s_contact,
    ]
    for i, fn in enumerate(builders, start=2):
        fn(prs, i)
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
